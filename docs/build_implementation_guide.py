"""Build a local Word implementation guide from current source, without loading secrets."""
from pathlib import Path
import json
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
doc = Document()
section = doc.sections[0]
section.top_margin = section.bottom_margin = Inches(0.7)
section.left_margin = section.right_margin = Inches(0.7)
doc.styles['Normal'].font.name = 'Calibri'
doc.styles['Normal'].font.size = Pt(10.5)
doc.styles['Normal'].paragraph_format.space_after = Pt(7)
for name, size in [('Title', 30), ('Heading 1', 19), ('Heading 2', 13)]:
    doc.styles[name].font.name = 'Calibri'
    doc.styles[name].font.size = Pt(size)
    doc.styles[name].font.color.rgb = RGBColor.from_string('126B70')
section.header.paragraphs[0].text = 'REQUIREMENT ANALYSIS AGENT | IMPLEMENTATION GUIDE'
footer = section.footer.paragraphs[0]
footer.add_run('Technical reference | Page ')
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
footer._p.append(field)
doc.core_properties.title = 'Requirement Analysis Agent - Complete Implementation and Flow'
doc.core_properties.author = 'Requirement Analysis Agent project'


def p(text):
    doc.add_paragraph(text)


def h(text, level=1):
    doc.add_heading(text, level)


def bullets(items):
    for text in items:
        doc.add_paragraph(text, 'List Bullet')


def steps(items):
    for text in items:
        doc.add_paragraph(text, 'List Number')


def code(text):
    paragraph = doc.add_paragraph()
    for i, line in enumerate(text.strip('\n').splitlines()):
        run = paragraph.add_run(('\n' if i else '') + line)
        run.font.name = 'Consolas'
        run.font.size = Pt(8)
    shade = OxmlElement('w:shd')
    shade.set(qn('w:fill'), 'F1F5F6')
    paragraph._p.get_or_add_pPr().append(shade)


def table(headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Light Shading Accent 1'
    for cell, value in zip(t.rows[0].cells, headers):
        cell.text = value
    for row in rows:
        for cell, value in zip(t.add_row().cells, row):
            cell.text = str(value)
    for row in t.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(9)
    doc.add_paragraph()


def chapter(title):
    doc.add_page_break()
    h(title)


doc.add_paragraph('PROJECT DOCUMENTATION', 'Subtitle')
doc.add_heading('Requirement Analysis Agent', 0)
doc.add_paragraph('Complete implementation, output structure and end-to-end flow', 'Subtitle')
p('Prepared: 27 October 2026 | Based on the current local source code')
p('This guide explains how a manually entered story or extracted document becomes a structured analysis through React, FastAPI and GEP. It covers input collection, prompts, JSON parsing, Pydantic validation, assessment scores, evidence checks, visual rendering, feedback and exports.')
h('The central idea', 2)
p('The AI generates the analysis content. Your application defines the field names, data types, validation rules and visual layout. The backend accepts a model response only after parsing and validation. A valid structure does not prove that the business interpretation is correct.')
h('Current upload behavior', 2)
p('Document upload is a separate section. Selecting a PDF or DOCX only extracts text for preview and editing. It does not automatically call GEP. The user clicks Analyze document to start analysis. Manual entry retains its own Analyze button.')
p('All examples are synthetic. This guide does not load or include actual API keys, environment-file contents or private user documents.')
h('Contents', 2)
bullets(['1. Architecture and ownership of output', '2. Source-code map', '3. Manual-entry flow', '4. Document-upload flow and extraction', '5. API contracts and schemas', '6. Prompt and GEP request', '7. Parsing and validation', '8. Story Assessment and score labels', '9. Evidence checks and response assembly', '10. Frontend state and rendering', '11. Reviews, feedback and exports', '12. Complete worked example', '13. Errors, timeouts and retries', '14. How to modify the implementation', '15. Local setup and testing', 'Appendices: current prompt, schema excerpts and glossary'])

chapter('1. Architecture and ownership of output')
table(['Layer', 'Responsibility'], [
    ('React + TypeScript', 'Collect input; show progress/errors; display findings; store local reviewer edits and feedback.'),
    ('FastAPI + Pydantic', 'Validate inputs; extract text; call the provider; validate model output; compute labels; check evidence; construct responses.'),
    ('GEP model', 'Generate finding content, initial classifications, numeric estimates, reasons and gaps from source fields.')])
h('End-to-end flow', 2)
steps(['Manual story -> click Analyze -> POST /api/v1/analyze.', 'PDF/DOCX selection -> POST /api/v1/documents/extract -> preview/edit text -> click Analyze document -> POST /api/v1/analyze.', 'Analysis route -> shared analysis service -> GEP provider -> system instructions plus user-input JSON.', 'Model JSON text -> parser -> Pydantic validation -> ID/evidence checks -> final API response.', 'API response -> React result state -> analysis cards and assessment -> review/feedback -> JSON or Markdown download.'])
h('Where the format comes from', 2)
p('The prompt describes the required model-output format. Backend schemas enforce the application data contract. React and export functions create the visible report format. Changing a prompt does not automatically change the API schema or screen layout.')
table(['Value', 'Source'], [('Finding title, description, quotes and initial classification', 'Model, followed by backend evidence checks.'), ('Complexity/readiness numbers, reasons and gaps', 'Model; backend validates types, presence and ranges.'), ('Complexity/readiness labels', 'Backend score lookup.'), ('Provider/model metadata, draft status and warnings', 'Backend application.'), ('Review decisions, edited finding content and usefulness feedback', 'Person using the frontend.')])

chapter('2. Source-code map')
table(['File', 'Purpose'], [
    ('backend/app/main.py', 'FastAPI creation, CORS, route registration and safe exception handlers.'),
    ('backend/app/api/routes.py', 'Health, provider listing, shared analysis route and bounded multipart extraction route.'),
    ('backend/app/core/config.py', 'Backend-only settings; SecretStr credential field.'),
    ('backend/app/schemas.py', 'Request, finding, analysis, assessment, final response and extracted-text contracts.'),
    ('backend/app/prompts/analysis_v1.txt', 'Instructions, score guides and complete JSON example for the model.'),
    ('backend/app/providers/base.py', 'Provider protocol and ProviderError contract.'),
    ('backend/app/providers/openai_provider.py', 'GEP configuration checks, Chat Completions request, status translation and one bounded retry.'),
    ('backend/app/services/model_output.py', 'Complete code-fence removal, JSON parsing and safe field-specific validation errors.'),
    ('backend/app/services/analyzer.py', 'Overall timeout, repeat validation, ID repair, evidence checks and final response assembly.'),
    ('backend/app/services/documents.py', 'In-memory PDF/DOCX extraction, table handling, signature/type checks and text limits.'),
    ('backend/app/assessment_labels.json', 'Shared score-to-label mapping imported by backend and frontend types.'),
    ('frontend/src/App.tsx', 'Provider/result/document state, manual submission, extraction, analysis, cancellation and retry.'),
    ('frontend/src/components/StoryForm.tsx', 'Manual form, Load example and Analyze.'),
    ('frontend/src/components/DocumentUpload.tsx', 'Separate document section, picker, filename, preview/editor and Analyze document.'),
    ('frontend/src/api/client.ts', 'Fetch helpers for JSON analysis, provider discovery and multipart extraction.'),
    ('frontend/src/types.ts', 'TypeScript contracts, section names and visible heading labels.'),
    ('frontend/src/components/AnalysisResults.tsx', 'Finding cards, review decisions, assessment and feedback.'),
    ('frontend/src/export.ts', 'Browser-generated JSON and Markdown reports.'),
    ('frontend/src/styles.css', 'Layout, cards, controls and responsive styles.'),
    ('backend/tests/ and frontend/src/test/', 'Offline backend, extraction, provider, browser and request-transport tests.')])

chapter('3. Manual-entry flow')
steps(['On mount, App requests GET /api/v1/providers. The response exposes provider ID, model and enabled flag, never the API key.', 'StoryForm owns the story, criteria, context and provider selection. Load example fills the manual form. Browser controls limit each text field to 12,000 characters.', 'Analyze rejects a blank story, then passes a Story object to App.submit().', 'submit() rejects duplicate in-progress submission, calls begin(), clears the previous result and document state, and stores a retry action.', 'runAnalysis() sets the analysis loading phase and calls analyzeStory(story, signal).', 'api/client.ts serializes the story as JSON and posts it to /api/v1/analyze.', 'FastAPI validates AnalysisRequest. The route finds the selected provider and calls the shared service.', 'The service requests a validated provider result under an overall timeout, checks IDs/evidence, and creates AnalysisResponse.', 'App stores a result only when its run ID matches the current request. AnalysisResults renders it. Only the current request can clear its busy state.'])
h('Illustrative input', 2)
request = {'user_story': 'As a registered customer, I want to reset my password using my verified email address so that I can regain access to my account.', 'acceptance_criteria': 'A reset link can only be used once.', 'context': 'The application already has an email service.', 'provider': 'openai'}
code(json.dumps(request, indent=2))
p('user_story is required and trimmed; whitespace-only input is rejected. acceptance_criteria and context default to empty strings. provider defaults to openai and is limited to 40 characters. The provider field selects the adapter and is excluded from the source data sent to the model.')

chapter('4. Document-upload flow and extraction')
steps(['DocumentPicker opens the browser selector. Upload can be used even before GEP is configured, because extraction is independent of model access.', 'App.upload() checks the active file identity and avoids submitting the same file twice during extraction. A new selection cancels the prior browser request and starts a new run.', 'Browser checks reject legacy .doc, unsupported extensions, empty files and files over 10 MB.', 'extractDocument() posts FormData with a file part to /api/v1/documents/extract. It lets the browser set the multipart Content-Type and boundary.', 'The route bounds the full multipart stream and accepts one file with no extra text form fields. Its spool threshold exceeds the bounded body size, keeping accepted files in memory.', 'The route reads at most the limit plus one byte, invokes extract_document through run_in_threadpool, and closes the multipart form/file resources afterward.', 'The extraction service checks MIME, extension and signature. PDF uses pypdf; DOCX uses python-docx. The response is plain text plus filename, with no model call.', 'App stores the extracted document. Upload success stops here. The user can expand Preview and edit extracted text.', 'Analyze document is enabled for nonblank text when the app is not busy and a provider is enabled. It calls rerun() and the existing analysis endpoint.', 'Document text is sent as user_story with empty acceptance_criteria and context. Manual-form criteria/context are not automatically combined with document text.'])
table(['Rule', 'Implementation'], [('File size', '10 x 1024 x 1024 bytes; UI calls this 10 MB.'), ('Multipart body', 'File limit plus 64 KiB for metadata; one file, zero extra fields.'), ('Extracted text', 'At most 12,000 characters; longer content is rejected, not truncated.'), ('MIME validation', 'application/pdf or DOCX Open XML MIME; application/octet-stream allowed with signature/parser checks.'), ('Signature', 'PDF starts %PDF-; DOCX starts PK ZIP signature and must include word/document.xml.'), ('PDF extraction', 'Strict reader; encrypted PDFs rejected; selectable page text joined with newlines.'), ('DOCX extraction', 'Body paragraphs/tables in order; cell separators; nested table traversal.'), ('Expanded DOCX size', 'Declared ZIP uncompressed total over 50 MiB rejected; encrypted ZIP entries rejected.'), ('No storage', 'Memory processing and closed upload resources; no database or retained upload.')])
p('PDF reading order and tables depend on the text layer; exact visual layout is not preserved. DOCX headers, footers, text boxes and embedded images are not explicitly traversed. OCR is not implemented.')
p('Exact no-selectable-text PDF message: This PDF does not contain selectable text. Please upload a text-based PDF or paste the user story manually.')

chapter('5. API contracts and schemas')
table(['Endpoint', 'Request', 'Response'], [('GET /health', 'None', 'Process health; does not check GEP.'), ('GET /api/v1/providers', 'None', 'Array of {id, model, enabled}.'), ('POST /api/v1/documents/extract', 'Multipart file field', '{filename, text}.'), ('POST /api/v1/analyze', 'AnalysisRequest JSON', 'AnalysisResponse JSON or safe error envelope.')])
code('''StrictModel (extra="forbid")
  AnalysisRequest
  Finding
  Analysis
    AssessedAnalysis -> inherits sections, adds story_assessment
  StoryAssessment -> scores, reasons, gaps; computed labels
  AnalysisResponse -> metadata + analysis + story_assessment
  DocumentTextResponse -> filename + text''')
table(['Contract', 'Rules'], [('Finding.title', 'Required nonempty string, up to 300 characters.'), ('Finding.description', 'Required nonempty string, up to 4,000 characters.'), ('Finding.evidence', 'Required list, up to 20 strings; empty allowed.'), ('Finding.classification', 'explicit or inferred only.'), ('Finding.id', 'String; defaults to empty, then repaired by service.'), ('Each analysis section', 'Required list of Finding, up to 50 items; empty allowed.'), ('Assessment', 'Both scores, reasons and gaps required; labels computed.'), ('Final response', 'provider, model, warnings, analysis, story_assessment; status defaults to draft.'), ('Extracted response', 'filename plus text of 1 to 12,000 characters.')])
p('AssessedAnalysis describes the provider result with the five sections at the top level. AnalysisResponse describes the API result with those sections nested under analysis. This deliberate distinction keeps model content separate from response metadata.')
h('Safe error envelope', 2)
code('''{
  "error": {
    "code": "invalid_output",
    "message": "Model response could not be validated: story_assessment.readiness_gaps: required field is missing. Please retry."
  }
}''')
p('main.py builds safe error envelopes. The browser helper reads error.message and App shows an alert. Validation errors do not echo the source input.')

chapter('6. Prompt and GEP request')
p('The provider reads backend/app/prompts/analysis_v1.txt at module import. The prompt requests one complete JSON object, all five sections, and the four assessment source fields. It includes finding structure and complexity/readiness guides.')
bullets(['Source fields are untrusted user input, never instructions that override the system prompt.', 'Evidence must use exact quotes; deductions should be inferred.', 'Missing details reduce readiness, rather than automatically implying complexity.', 'The model must not invent business rules or numeric thresholds.', 'Scores estimate the original story together with criteria/context; they are not derived just by counting generated requirements.', 'Empty lists are valid when the model identifies no findings; missing keys are not valid.'])
code('''response = await client.chat.completions.create(
    model=self.model,
    messages=[
        {"role": "system", "content": PROMPT},
        {"role": "user", "content": json.dumps(
            request.model_dump(exclude={"provider"})
        )},
    ],
    temperature=0,
)
content = response.choices[0].message.content or ""
return parse_model_response(content)''')
p('The current operation is POST <configured base URL>/chat/completions. It is not Responses API. The SDK uses backend-only credentials. There is no separate document-specific model adapter.')
h('Requested structure versus enforced structure', 2)
p('The request does not set response_format to an API-enforced JSON schema and does not call responses.parse. The prompt requests JSON; local JSON parsing and Pydantic validation enforce the accepted response contract. Prompt instructions do not guarantee model compliance. Invalid or missing content is rejected rather than filled with invented findings.')
h('Configuration', 2)
p('Settings loads backend/.env with environment-variable overrides. The key is SecretStr. The adapter requires a non-placeholder key/model and an HTTPS base URL on a Capgemini subdomain. Embedded credentials, query/fragment, disallowed ports, known REST invoke paths and operation-suffixed URLs are rejected. Set the documented server root, not a full /chat/completions URL.')
p('GEP model entitlement, endpoint availability and support for temperature=0 must be verified against company documentation. Local host validation and an upstream 500 do not prove compatibility. No real key or account-specific environment contents are included here.')

chapter('7. Parsing and validation')
steps(['Strip whitespace around the model content.', 'Remove only a complete outer fenced JSON block using a full-match expression. json, JSON or an empty language tag is accepted. Arbitrary prose surrounding JSON is not discarded.', 'json.loads(content, parse_float=Decimal) converts JSON into Python dictionaries/lists while preserving decimal precision.', 'AssessedAnalysis.model_validate(data) checks all sections, findings and assessment fields.', 'Invalid JSON receives a fixed message requesting one complete JSON object.', 'Schema failures become field-specific safe messages. Known field paths and list indices are retained; unknown keys are masked. Values, validator context and raw model text are not copied into diagnostics. Up to ten issues are shown.'])
code('''data = json.loads(content, parse_float=Decimal)
result = AssessedAnalysis.model_validate(data)''')
p('JSON parsing asks whether the text is valid JSON. Schema validation asks whether the resulting data matches the application contract. TypeScript declarations do not perform these backend checks.')
table(['Problem', 'Outcome'], [('Prose or malformed JSON', 'Parsing failure; no report invented.'), ('Missing risks', 'Required-field error naming risks.'), ('Missing readiness_gaps', 'Required-field error even though an explicitly supplied [] is valid.'), ('classification: approved', 'Rejected; explicit/inferred only.'), ('Unexpected key', 'Rejected by extra=forbid; diagnostic masks unknown key.'), ('Score "3" or 3.0', 'Normalized to integer 3 if safe.'), ('Score true, 2.5, 0 or 6', 'Rejected.'), ('Missing/mismatched model label', 'Label ignored or absent; backend derives it from score.')])
p('Required analysis and assessment content is never invented to satisfy validation. A finding ID is a documented exception: it may default to an empty string and then be repaired deterministically. Optional request fields have their own documented defaults.')

chapter('8. Story Assessment and score labels')
labels = json.loads((ROOT / 'backend/app/assessment_labels.json').read_text())
complexity = ['One simple action, few or no rules.', 'Simple workflow with a small number of rules.', 'Several steps or moderate business rules.', 'Multiple rules, roles, integrations or error cases.', 'Broad workflow with dependencies, conditional paths or security/compliance concerns.']
readiness = ['Goal or user is unclear.', 'Goal understandable but major details missing.', 'Basic story clear but important questions remain.', 'Mostly clear/testable, with minor gaps.', 'Clear, specific and testable.']
table(['Score', 'Complexity', 'Readiness'], [(str(i), labels['complexity'][str(i)] + ': ' + complexity[i-1], labels['readiness'][str(i)] + ': ' + readiness[i-1]) for i in range(1, 6)])
p('Both scores ultimately become strict integers in the range 1 to 5. The before-validator rejects booleans, trims strings, limits string length to 32 and checks simple signed-decimal syntax. String/float/Decimal values must be finite, integral and within range before conversion. Fractional scores are not rounded. Scientific-notation strings are not accepted by the string syntax rule.')
p('Decimal JSON parsing prevents a fractional value such as 3.00000000000000001 from rounding to 3 before validation. Numeric strings such as "3.0" are accepted because they represent a whole number.')
p('complexity_reasons contains 1 to 10 nonblank strings. readiness_gaps contains 0 to 20 nonblank strings. Each entry is at most 1,000 characters. An empty gap list means the AI identified no gaps, not that the story is guaranteed complete.')
p('ignore_model_labels removes model-supplied label fields. Computed properties derive labels from backend/app/assessment_labels.json. Frontend types import the same mapping. The API returns the computed labels and the UI displays them.')
code('''@computed_field
@property
def complexity_label(self) -> str:
    return ASSESSMENT_LABELS["complexity"][str(self.complexity_score)]''')
p('Scores and explanations are AI estimates requiring human review. The code validates ranges and structure; it does not turn those estimates into objective measurements.')

chapter('9. Evidence checks and response assembly')
p('services/analyzer.py calls provider.analyze under asyncio.wait_for and validates the returned object again. This protects the service boundary if another provider returns a plain dictionary or invalid data.')
h('ID repair', 2)
p('A valid ID starts with a letter and has at most 64 characters consisting of letters, digits, underscore or hyphen. Valid IDs are reserved first. Invalid/repeated IDs are repaired with FR, NFR, RISK, ASM or Q prefixes and numbered suffixes. -R is appended when needed to avoid collisions. Processing the same ordered result is deterministic; independent model runs may differ.')
h('Evidence checks', 2)
p('A retained quote must be nonblank and an exact substring of user_story or acceptance_criteria. A quote found only in context is not retained under this rule. An explicit finding with no supported quotes, or with any unsupported quote, becomes inferred and receives a warning. Unsupported quotes are removed with another warning.')
p('For example, "reset my password" may be retained when present verbatim. "Link expires in 15 minutes" is removed if absent from the story and criteria. These checks provide traceability; they do not prove logical correctness or eliminate hallucinations.')
h('Build the final API response', 2)
code('''return AnalysisResponse(
    provider=provider.name,
    model=provider.model,
    warnings=warnings,
    analysis=Analysis.model_validate(
        result.model_dump(exclude={"story_assessment"})
    ),
    story_assessment=result.story_assessment,
)''')
p('The five model sections become the nested analysis object. Assessment stays beside analysis. Backend code adds metadata/warnings; status defaults to draft. The route declares response_model=AnalysisResponse and FastAPI serializes this contract.')

chapter('10. Frontend state and rendering')
p('types.ts defines section keys and display headings. AnalysisResults loops through sections and reads result.analysis[section]. Each finding becomes a card with ID, badge, title/description inputs, evidence and a review control. Empty sections show a no-findings message; warnings appear in expandable notes.')
p('Story Assessment displays both scores, labels, reasons and gaps with an AI-estimate notice. The frontend type allows absent/null assessment for older results, showing an unavailable message. The current backend requires assessment content. TypeScript catches development-time shape errors; it does not validate arbitrary network JSON at runtime.')
table(['State/reference', 'Purpose'], [('providers / chosenProvider', 'Available adapters and selected provider.'), ('result', 'Current report including reviewer edits and submitted feedback.'), ('document', 'Filename and extracted/editable text.'), ('busy / phase / error', 'Shared activity, reading/analysis state and safe message.'), ('current', 'Incrementing run ID; stale completions cannot update state.'), ('controller', 'AbortController for the current browser request.'), ('inProgress', 'Synchronous duplicate-submission guard.'), ('activeFile', 'Identity of the active extraction; avoids duplicate same-file reads.'), ('retryAction', 'Repeats the failed extraction or analysis stage.')])
h('Race protection', 2)
p('begin() aborts the prior request, increments the run ID, clears old errors/results and marks busy. Each continuation checks its captured run ID before updating the UI. Only the latest request can clear busy in finally. Cleanup aborts and invalidates active work on unmount. This protects against work that finishes even after browser cancellation.')
p('Cancellation does not guarantee that already-running backend/model work stops. Its response is ignored when stale. A new file may replace an ongoing run. The preview editor and analysis buttons are disabled while busy. Starting manual analysis clears document state.')
p('The HTTP helper combines a caller abort signal with a 190-second timeout. It uses JSON for analysis and FormData for extraction. Document analysis uses the provider selected in the manual form but passes empty criteria/context. A new file is extracted only; it requires another explicit analysis click.')

chapter('11. Reviews, feedback and exports')
p('Reviewer annotations live in React memory. Editing a title/description updates the corresponding finding with an immutable state update. Review can be pending, accepted or rejected. These are human decisions; the agent never automatically approves requirements.')
p('Feedback asks Was this analysis useful? with a 1-5 rating and optional comment. Submission adds feedback: {rating, comment} to the report. Comments are trimmed and the editor limits them to 2,000 characters. The selected submitted rating is shown. There is no database or sessionStorage persistence; refresh or a new run clears the current draft.')
table(['Export', 'How it is created'], [('JSON', 'JSON.stringify(result, null, 2); includes edits, review decisions, assessment and submitted feedback.'), ('Markdown', 'markdown(result) constructs headings, findings, evidence quotes, decisions, assessment and feedback.'), ('Download mechanism', 'Browser Blob -> object URL -> anchor click -> URL revocation.')])
p('The model does not create the React layout or export styling. AnalysisResults.tsx renders structured data; styles.css supplies visual rules. export.ts independently formats the report. The current application downloads JSON/Markdown, not Word; this implementation guide is a separate Word artifact.')
p('Exports serialize the result object. They do not automatically include the source file, filename or extracted source text because App stores those separately in document state. Save a report before refreshing if reviewer decisions must be preserved.')

chapter('12. Complete worked example')
p('All values here are synthetic. Empty arrays demonstrate structure and do not assert what a real model would find for this story.')
h('A. Browser request', 2)
code(json.dumps(request, indent=2))
model = {name: [] for name in ['functional_requirements', 'non_functional_requirements', 'risks', 'assumptions', 'clarification_questions']}
model['functional_requirements'] = [{'id': 'FR-001', 'title': 'Request password reset', 'description': 'Allow the registered customer to request a reset using their verified email address.', 'evidence': ['reset my password using my verified email address'], 'classification': 'explicit'}]
model['story_assessment'] = {'complexity_score': '2.0', 'complexity_reasons': ['A single reset workflow with a small number of stated rules.'], 'readiness_score': 3, 'readiness_gaps': ['Clarify link expiry and behavior for an unrecognized email address.']}
h('B. Illustrative model response', 2)
code(json.dumps(model, indent=2))
p('The string "2.0" demonstrates safe conversion. Model labels are not needed.')
chapter('12. Worked example: final response and annotations')
final = {'provider': 'openai', 'model': 'configured-model-id', 'status': 'draft', 'warnings': [], 'analysis': {k: v for k, v in model.items() if k != 'story_assessment'}, 'story_assessment': {**model['story_assessment'], 'complexity_score': 2, 'complexity_label': 'Low', 'readiness_label': 'Needs clarification'}}
code(json.dumps(final, indent=2))
p('Backend processing normalizes 2.0 to integer 2, derives labels, retains the supported quote and moves the sections under analysis. It does not invent the description or the numeric estimates. warnings is empty here because the example evidence and ID pass their checks.')
h('Human annotations added later', 2)
code('''review: "accepted"
feedback: {
  rating: 4,
  comment: "Useful starting point; clarify expiry."
}''')
p('review belongs on the reviewed finding; feedback belongs on the report. They are added in frontend state after the backend response and appear in subsequent downloads.')

chapter('13. Errors, timeouts and retries')
table(['Condition', 'Outcome'], [('Unknown provider', '400 unknown_provider.'), ('Invalid story/request', '422 invalid_input; fixed safe message.'), ('Unsupported extension/MIME', '415; .doc has specific save-as-DOCX guidance.'), ('Malformed multipart/missing file', '400 invalid_upload.'), ('Oversized file', '413 document_too_large.'), ('Corrupt/protected/empty/textless document', '422 with safe error; no raw parser exception.'), ('Extracted text too long', '422 document_text_too_long; no truncation.'), ('Invalid/missing model content', 'Normally 502 invalid_output; sanitized field-specific message.'), ('Authentication or permission failure', 'Mapped to authentication_error/access_denied; 503 to client; no retry.'), ('Unsupported model/endpoint/parameter', 'Specific safe configuration guidance; no retry.'), ('Quota/rate limit', '429; no custom automatic retry.'), ('Eligible upstream 500/502/503/504', 'One retry after 0.5 seconds; persistent failure returns 503 gateway_failure.'), ('Transport connection failure', 'gateway_connection_error; no custom retry.'), ('SDK timeout or service deadline', '504 timeout.'), ('Unexpected backend failure', '500 with fixed generic message.')])
p('SDK max_retries is 0. The application uses at most two attempts. Retry requires the first attempt to have status 500, 502, 503 or 504 and not carry a known non-retryable sanitized code. Authentication, permission, unsupported model/endpoint/parameter, request, quota and rate-limit codes prevent retry even with a 5xx status. Unknown upstream codes are masked as unrecognized_upstream_error and remain eligible when paired with a transient status.')
p('Logs retain the numeric upstream status, an allowlisted code and a fixed application message for each failed attempt. Arbitrary upstream messages, headers, keys, source documents and raw model content are not copied into these logs. The PDF parser logger is silenced because its warnings may contain document fragments.')
p('Persistent eligible transient failure is shown as: GEP is temporarily unavailable; please retry. This does not identify the root cause. Check the documented GEP server, model entitlement, operation and supported parameters when investigating.')
p('The service deadline defaults to 60 seconds and may be configured from 1 to 180. It includes both attempts and the delay. Browser requests time out after 190 seconds. UI Retry repeats extraction after a read failure; after analysis failure it resends captured text without another extraction. Successful extraction still requires Analyze document.')

chapter('14. How to modify the implementation')
table(['Desired change', 'Files'], [('Change model instructions', 'prompts/analysis_v1.txt and prompt-related tests.'), ('Add/change required output field', 'Prompt, schemas.py, frontend types.ts, rendering/export and tests.'), ('Change score-label mapping', 'assessment_labels.json and score tests; verify frontend consumers.'), ('Rename a visible section', 'labels in frontend/src/types.ts; API key can remain unchanged.'), ('Change output card layout/colors', 'AnalysisResults.tsx and styles.css.'), ('Change Markdown report design', 'frontend/src/export.ts.'), ('Change extraction', 'documents.py, extraction route and test_documents.py.'), ('Change upload/analysis trigger', 'App.upload()/rerun() and DocumentUpload.tsx.'), ('Change GEP operation/parameters', 'openai_provider.py and mocked provider tests.'), ('Add provider', 'Implement Provider protocol; register in get_providers; add offline tests.')])
h('Example future change: priority on each finding', 2)
steps(['Define allowed values and whether priority is required. It is not a current field.', 'Add priority to Finding with an explicit validation rule.', 'Update the prompt so every finding includes priority.', 'Update the TypeScript Finding type.', 'Display it in AnalysisResults and add it to Markdown output. JSON serializes fields that are in the result object.', 'Add valid/missing/invalid tests, run backend/frontend checks, then restart the backend.'])
p('Do not add an invented default just to conceal missing model content. Keep backend, prompt, frontend and exports aligned whenever changing a required field.')
h('How the implementation is composed', 2)
p('The shared GEP adapter and response schema handle both input sources. Assessment extends the result contract. Safe parsing handles fences, required fields and whole-number scores. The service computes labels and checks evidence. Gateway logic handles one bounded transient retry. A separate extraction endpoint converts documents to plain text. The current document UI intentionally separates extraction from model analysis. No second LLM integration, OCR or database was introduced.')

chapter('15. Local setup and testing')
p('Use Windows PowerShell. Start each terminal at the repository root. Configure backend/.env locally before a live model call; never copy the key into frontend code or generated reports.')
h('Backend terminal', 2)
code('''cd .\\backend
.\\.venv\\Scripts\\python.exe -m pip install -r requirements.txt
.\\.venv\\Scripts\\python.exe -m uvicorn app.main:app --reload --host localhost --port 8000''')
h('Frontend terminal', 2)
code('''cd .\\frontend
npm.cmd install
.\\start.cmd''')
p('Open http://localhost:5173. Backend API docs: http://localhost:8000/docs. Health: http://localhost:8000/health. Virtual environment activation is unnecessary. If a copied .venv points to a missing interpreter, recreate it with a locally installed compatible Python and reinstall dependencies.')
h('Try the complete flow', 2)
steps(['Manual: use Load example or enter a story; add optional criteria/context; click Analyze.', 'Document: choose a short text-based PDF/DOCX in Document analysis; wait for extraction. Confirm that analysis does not start automatically.', 'Expand the preview, correct text if needed, then click Analyze document.', 'Review the five sections, warnings and assessment. Edit a finding and mark its review decision.', 'Submit usefulness feedback and download JSON/Markdown. Check annotations in the file.', 'Select another file during a run and confirm older requests cannot replace the newest text or result.'])
h('Test commands', 2)
code('''# In backend
.\\.venv\\Scripts\\python.exe -m pytest -q

# In frontend
npm.cmd test
npm.cmd run build''')
p('Recorded session verification: 136 backend tests passed after extraction implementation. Following the latest explicit Analyze document UI change, 20 frontend tests and the production build passed. These are previous recorded checks, not a new execution performed just to create this guide. Provider tests use mocks; live GEP entitlement/connectivity was not verified.')
p('Backend coverage includes schemas, scores/labels, parsing/fences, missing fields, IDs/evidence, timeouts, retries/no-retry cases, safe diagnostics and extraction of PDFs/DOCX paragraphs/tables. Upload tests include signatures, size/text limits, corrupt/protected/empty/scanned documents and memory lifecycle. Frontend tests cover manual entry, extraction without auto-analysis, explicit document analysis, preview edits, retries, duplicate prevention, stale responses, feedback, exports and multipart transport.')

chapter('Appendix A. Current model prompt')
p('Verbatim source: backend/app/prompts/analysis_v1.txt. This contains application instructions and synthetic examples only.')
code((ROOT / 'backend/app/prompts/analysis_v1.txt').read_text())
chapter('Appendix B. Backend contract excerpts')
p('Verbatim excerpts from backend/app/schemas.py. Score validation and computed labels are explained in section 8.')
schema = (ROOT / 'backend/app/schemas.py').read_text()
code(schema[schema.index('class Finding'):schema.index('ASSESSMENT_LABELS')])
code(schema[schema.index('class AssessedAnalysis'):])
p('functional_requirements is a field key. Functional requirements is a visible heading. result.analysis.functional_requirements is the path used to access that array in a final response. Renaming the heading alone does not rename the field.')
h('Glossary', 2)
table(['Term', 'Meaning'], [('JSON', 'Text format representing structured objects, arrays and primitive values.'), ('Schema', 'Contract for keys, types, allowed values and limits.'), ('Pydantic model', 'Python class representing and validating data.'), ('Prompt', 'Instructions asking the model for analysis in a particular shape.'), ('System message', 'Application instructions above the source-data user message.'), ('Parsing', 'Convert JSON text into Python data.'), ('Validation', 'Check that data against schema and rules.'), ('Serialization', 'Convert objects into JSON or report text.'), ('Explicit', 'Claimed direct support from source text.'), ('Inferred', 'Deduction or finding whose explicit evidence was not confirmed.'), ('Run ID', 'Incrementing identity used to reject stale async updates.'), ('Draft', 'Result awaiting human review, not automatic approval.')])

if __name__ == '__main__':
    path = ROOT / 'docs/Requirement_Analysis_Agent_Implementation_Guide.docx'
    doc.save(path)
    print(path)
    print(f'Paragraphs: {len(doc.paragraphs)}; tables: {len(doc.tables)}')
