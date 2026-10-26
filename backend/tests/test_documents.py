from io import BytesIO
import pytest
from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
from fastapi.testclient import TestClient
from app.main import app
from app.services.documents import MAX_FILE_SIZE, SCANNED_PDF_MESSAGE

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

@pytest.fixture
def client():
    with TestClient(app) as client:
        yield client

def pdf(text=None, encrypted=False):
    writer=PdfWriter(); page=writer.add_blank_page(width=612,height=792)
    if text:
        font=DictionaryObject({NameObject('/Type'):NameObject('/Font'),NameObject('/Subtype'):NameObject('/Type1'),NameObject('/BaseFont'):NameObject('/Helvetica')})
        page[NameObject('/Resources')]=DictionaryObject({NameObject('/Font'):DictionaryObject({NameObject('/F1'):writer._add_object(font)})})
        stream=DecodedStreamObject();stream.set_data(('BT /F1 12 Tf 72 720 Td ('+text+') Tj ET').encode())
        page[NameObject('/Contents')]=writer._add_object(stream)
    if encrypted: writer.encrypt('password')
    output=BytesIO();writer.write(output);return output.getvalue()

def docx(text='',table=False):
    document=Document()
    if text: document.add_paragraph(text)
    if table:
        row=document.add_table(rows=1,cols=2).rows[0]
        row.cells[0].text='Role'; row.cells[1].text='Registered customer'
        document.add_paragraph('After the table')
    output=BytesIO();document.save(output);return output.getvalue()

def upload(client,data,name='story.pdf',mime='application/pdf'):
    return client.post('/api/v1/documents/extract',files={'file':(name,data,mime)})

def test_pdf_extracts_text_without_llm(client):
    response=upload(client,pdf('As a customer I want to reset my password'))
    assert response.status_code==200
    assert response.json()['filename']=='story.pdf'
    assert 'reset my password' in response.json()['text']

def test_docx_extracts_paragraphs_and_tables_in_order(client):
    response=upload(client,docx('Reset password',True),'story.docx',DOCX_MIME)
    assert response.status_code==200
    assert response.json()['text']=='Reset password\nRole | Registered customer\nAfter the table'

@pytest.mark.parametrize('name,mime', [('story.txt','text/plain'),('story.doc','application/msword'),('story.pdf',DOCX_MIME),('story.docx','application/pdf')])
def test_unsupported_type(client,name,mime):
    assert upload(client,b'content',name,mime).status_code==415

@pytest.mark.parametrize('data,name,mime', [(b'', 'empty.pdf','application/pdf'),(b'%PDF-broken','broken.pdf','application/pdf'),(b'PK\x03\x04broken','broken.docx',DOCX_MIME),(b'not a PDF','fake.pdf','application/pdf')])
def test_empty_corrupt_and_signature_errors(client,data,name,mime):
    response=upload(client,data,name,mime)
    assert response.status_code==422
    assert response.json()['error']['code'] in {'empty_document','corrupt_document'}

def test_oversized_file(client):
    response=upload(client,b'%PDF-'+b'x'*MAX_FILE_SIZE)
    assert response.status_code==413
    assert '10 MB' in response.json()['error']['message']

def test_oversized_multipart_stream(client):
    response=upload(client,b'x'*(MAX_FILE_SIZE+65536))
    assert response.status_code==413

def test_encrypted_pdf(client):
    response=upload(client,pdf('private',True))
    assert response.status_code==422
    assert response.json()['error']['code']=='protected_document'

@pytest.mark.parametrize('data', [pdf(),pdf('   ')])
def test_scanned_or_blank_pdf_message(client,data):
    response=upload(client,data)
    assert response.status_code==422
    assert response.json()['error']['message']==SCANNED_PDF_MESSAGE

def test_empty_docx(client):
    response=upload(client,docx(),'empty.docx',DOCX_MIME)
    assert response.status_code==422
    assert response.json()['error']['code']=='no_extractable_text'

def test_encrypted_word_container(client):
    response=upload(client,bytes.fromhex('D0CF11E0A1B11AE1')+b'content','protected.docx',DOCX_MIME)
    assert response.status_code==422
    assert response.json()['error']['code']=='protected_document'

def test_extracted_text_limit(client):
    response=upload(client,docx('x'*12001),'long.docx',DOCX_MIME)
    assert response.status_code==422
    assert response.json()['error']['code']=='document_text_too_long'

def test_file_limit_and_missing_field(client):
    assert client.post('/api/v1/documents/extract',files={'other':('story.pdf',pdf(),'application/pdf')}).status_code==400
    assert client.post('/api/v1/documents/extract',files=[('file',('a.pdf',pdf(),'application/pdf')),('file',('b.pdf',pdf(),'application/pdf'))]).status_code==400

def test_content_is_never_logged(client,caplog):
    response=upload(client,docx('private-document-secret'),'story.docx',DOCX_MIME)
    assert response.status_code==200
    assert 'private-document-secret' not in caplog.text

def test_upload_stays_in_memory_and_is_closed(client,monkeypatch):
    import starlette.formparsers
    from tempfile import SpooledTemporaryFile
    from unittest.mock import patch
    files=[]
    def memory_file(*args,**kwargs):
        file=SpooledTemporaryFile(*args,**kwargs);files.append(file);return file
    monkeypatch.setattr(starlette.formparsers,'SpooledTemporaryFile',memory_file)
    with patch('app.api.routes.extract_document',return_value='Valid story'):
        response=upload(client,b'%PDF-'+b'x'*(MAX_FILE_SIZE-5))
    assert response.status_code==200
    assert files and all(not file._rolled and file.closed for file in files)

def test_missing_boundary_and_wrong_request_type(client):
    response=client.post('/api/v1/documents/extract',content=b'invalid',headers={'content-type':'multipart/form-data'})
    assert response.status_code==400
    assert client.post('/api/v1/documents/extract',json={}).status_code==415

def test_image_only_scanned_pdf(client):
    from pypdf.generic import NumberObject
    writer=PdfWriter();page=writer.add_blank_page(width=612,height=792)
    image=DecodedStreamObject();image.set_data(b'\x00\x00\x00')
    image.update({NameObject('/Type'):NameObject('/XObject'),NameObject('/Subtype'):NameObject('/Image'),NameObject('/Width'):NumberObject(1),NameObject('/Height'):NumberObject(1),NameObject('/ColorSpace'):NameObject('/DeviceRGB'),NameObject('/BitsPerComponent'):NumberObject(8)})
    page[NameObject('/Resources')]=DictionaryObject({NameObject('/XObject'):DictionaryObject({NameObject('/Scan'):writer._add_object(image)})})
    stream=DecodedStreamObject();stream.set_data(b'q 500 0 0 700 20 20 cm /Scan Do Q')
    page[NameObject('/Contents')]=writer._add_object(stream)
    data=BytesIO();writer.write(data)
    response=upload(client,data.getvalue())
    assert response.status_code==422
    assert response.json()['error']['message']==SCANNED_PDF_MESSAGE
